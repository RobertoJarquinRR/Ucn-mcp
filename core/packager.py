"""Packager: turn plain text into a submittable file (PDF / DOCX).

Moodle assignments here accept files only, so a text answer submitted as a file
is wrapped into a PDF (or DOCX when the assignment asks for Word).
"""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
ARTIFACTS = PROJECT_ROOT / "storage" / "artifacts"

# Unicode-capable fonts shipped with Windows
_FONT_CANDIDATES = [
    r"C:\Windows\Fonts\arial.ttf",
    r"C:\Windows\Fonts\calibri.ttf",
    r"C:\Windows\Fonts\segoeui.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
]


def _find_font() -> str | None:
    for f in _FONT_CANDIDATES:
        if Path(f).exists():
            return f
    return None


def _safe_name(s: str) -> str:
    return "".join(c if (c.isalnum() or c in " -_") else "_" for c in s).strip()[:60]


def text_to_pdf(text: str, title: str = "", author: str = "") -> Path:
    """Render `text` into a PDF and return its path."""
    from fpdf import FPDF

    ARTIFACTS.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = ARTIFACTS / f"{_safe_name(title) or 'trabajo'}_{ts}.pdf"

    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()

    font = _find_font()
    if font:
        pdf.add_font("main", "", font)
        pdf.set_font("main", size=12)
    else:
        pdf.set_font("Helvetica", size=12)

    if title:
        pdf.set_font_size(16)
        if font:
            pdf.set_font("main", size=16)
        else:
            pdf.set_font("Helvetica", "B", size=16)
        pdf.multi_cell(0, 10, title)
        pdf.ln(2)
        if font:
            pdf.set_font("main", size=12)
        else:
            pdf.set_font("Helvetica", size=12)

    if author:
        pdf.set_font_size(10)
        pdf.multi_cell(0, 6, author)
        pdf.ln(3)

    # Normalize the text: strip markdown-ish noise, keep paragraphs
    clean = text.replace("\r\n", "\n")
    clean = re.sub(r"^[#>*\-]+\s?", "", clean, flags=re.MULTILINE)
    for para in clean.split("\n"):
        if para.strip():
            pdf.multi_cell(0, 7, para)
        else:
            pdf.ln(4)

    pdf.output(str(out))
    return out


def text_to_docx(text: str, title: str = "") -> Path:
    """Render `text` into a .docx (requires python-docx)."""
    from docx import Document

    ARTIFACTS.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = ARTIFACTS / f"{_safe_name(title) or 'trabajo'}_{ts}.docx"

    doc = Document()
    if title:
        doc.add_heading(title, level=1)
    for para in text.replace("\r\n", "\n").split("\n"):
        doc.add_paragraph(para)
    doc.save(str(out))
    return out


def package_text(text: str, title: str = "", author: str = "",
                 fmt: str = "pdf") -> Path:
    fmt = (fmt or "pdf").lower()
    if fmt in ("doc", "docx", "word"):
        return text_to_docx(text, title)
    return text_to_pdf(text, title, author)
